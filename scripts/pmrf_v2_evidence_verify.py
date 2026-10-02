#!/usr/bin/env python3
"""
PMRF v2.0 — Phase 2-6: Evidence Verification Audit
Verifies PMRF v1.5 material claims against actual runtime artifacts.
Output: research/pmrf_v2_evidence_audit_20261002.json
"""
import json, datetime, hashlib, re
from pathlib import Path

ROOT = Path("/home/z/my-project")
OUT = ROOT / "research" / "pmrf_v2_evidence_audit_20261002.json"
TZ = datetime.timezone(datetime.timedelta(hours=3))  # Asia/Aden UTC+3
TS = datetime.datetime.now(TZ).isoformat(timespec="seconds")

report = {
    "task": "PMRF v2.0 evidence verification audit (v1.5 claims vs runtime artifacts)",
    "execution_timestamp": TS,
    "artifacts": {},
    "claim_checks": [],
    "runtime_counts": {},
}


def stat(rel):
    p = ROOT / rel
    if not p.exists():
        return {"exists": False}
    s = p.stat()
    h = hashlib.sha256(p.read_bytes()).hexdigest()[:16]
    return {
        "exists": True,
        "bytes": s.st_size,
        "mtime": datetime.datetime.fromtimestamp(s.st_mtime, TZ).isoformat(timespec="seconds"),
        "sha256_16": h,
    }


def check(claim, expected, condition, detail):
    report["claim_checks"].append({
        "claim": claim,
        "expected": expected,
        "result": "PASS" if condition else "FAIL",
        "detail": detail,
    })


# ---------- 1. Key artifact inventory ----------
KEY_ARTIFACTS = [
    "download/PROJECT MASTER REFERENCE FILE.md",
    "worklog.md",
    "research/sv_master_refresh_20261002.json",
    "research/sv_master_gateway_pricing_20261002.json",
    "research/sv_master_catalog_consolidated_20261002.json",
    "download/جدول_الكتالوج_الموحد_2026-10-02.xlsx",
    "download/مصفوفة_حملة_التفاوض_2026-10-02.xlsx",
    "download/دليل_حملة_التفاوض_الافتتاحية_2026-10-02.docx",
    "research/g1_cb_findings_20261002.json",
    "research/g1_cb_live_catalog_20261002.json",
    "research/g1_rescan_analysis_20261002.json",
    "research/g1_r3_newkeys_analysis_20261002.json",
    "research/g1_r3_final_20261002.json",
    "research/g1a_r5_decohomz_origin_20261002.json",
    "research/g1a_r6_websearch_sookbit_20261002.json",
    "download/G1a_تقرير_هوية_مشغل_المنصة_2026-10-02.md",
    "download/G1-R3_تقرير_جولة_المفاتيح_2026-10-02.md",
    "download/G1-RESCAN-R2_تقرير_إعادة_المسح_2026-10-02.md",
]
for k in KEY_ARTIFACTS:
    report["artifacts"][k] = stat(k)

# ---------- 2. Runtime counts (fresh, for §4 Runtime Inventory) ----------
def count_files(pattern_root):
    p = ROOT / pattern_root
    if not p.exists():
        return 0
    n = 0
    for f in p.rglob("*"):
        if f.is_file():
            n += 1
    return n

report["runtime_counts"] = {
    "research_files": count_files("research"),
    "scripts_files": count_files("scripts"),
    "download_files": count_files("download"),
    "notion_raw_files": count_files("notion_raw"),
    "data_research_files": count_files("data/research"),
    "worklog_lines": len((ROOT / "worklog.md").read_text(encoding="utf-8").splitlines()),
    "worklog_task_ids": len(re.findall(r"^Task ID:", (ROOT / "worklog.md").read_text(encoding="utf-8"), re.M)),
    "pmrf_lines": len((ROOT / "download/PROJECT MASTER REFERENCE FILE.md").read_text(encoding="utf-8").splitlines()),
    "pmrf_sections_h2": len(re.findall(r"^## ", (ROOT / "download/PROJECT MASTER REFERENCE FILE.md").read_text(encoding="utf-8"), re.M)),
}

# ---------- 3. Claim: consolidated catalog == 1,112 records ----------
try:
    cons = json.loads((ROOT / "research/sv_master_catalog_consolidated_20261002.json").read_text())
    entities = cons.get("entities", {})
    gateway = cons.get("gateway", {})
    total = 0
    per_entity = {}
    for name, blk in entities.items():
        if isinstance(blk, dict):
            for k in ("products", "items", "catalog", "records"):
                if isinstance(blk.get(k), list):
                    per_entity[name] = len(blk[k])
                    total += len(blk[k])
                    break
            else:
                # maybe the block itself is a list of records
                per_entity[name] = "STRUCT?"
    gw = 0
    for k in ("products", "items", "models"):
        if isinstance(gateway.get(k), list):
            gw = len(gateway[k])
            break
    else:
        gw = gateway.get("count", "?")
    check(
        "الكتالوج الموحد = 1,112 سجلًا (11 كيانًا + البوابة)",
        1112,
        total + (gw if isinstance(gw, int) else 0) == 1112,
        {"per_entity": per_entity, "entity_sum": total, "gateway": gw, "grand_total": total + (gw if isinstance(gw, int) else 0)},
    )
except Exception as e:
    check("الكتالوج الموحد = 1,112", 1112, False, f"EXC {e}")

# ---------- 4. Claim: neg campaign xlsx = 8 sheets, 30 scenarios ----------
try:
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "download/مصفوفة_حملة_التفاوض_2026-10-02.xlsx", read_only=True, data_only=False)
    sheets = wb.sheetnames
    sc_count = None
    if "السيناريوهات" in sheets or any("سيناريو" in s for s in sheets):
        sname = next(s for s in sheets if "سيناريو" in s)
        ws = wb[sname]
        ids = 0
        for row in ws.iter_rows(min_col=1, max_col=2):
            v = row[0].value
            if isinstance(v, str) and re.match(r"^S-[A-Z]\d+$", v.strip()):
                ids += 1
        sc_count = ids
    check("مصفوفة حملة التفاوض = 8 أوراق", 8, len(sheets) == 8, {"sheets": sheets})
    check("حملة التفاوض = 30 سيناريو (S-x y)", 30, sc_count == 30, {"scenario_id_count": sc_count})
    wb.close()
except Exception as e:
    check("مصفوفة حملة التفاوض 8 أوراق/30 سيناريو", "8/30", False, f"EXC {e}")

# ---------- 5. Claim: unified catalog xlsx = 7 sheets ----------
try:
    from openpyxl import load_workbook
    wb2 = load_workbook(ROOT / "download/جدول_الكتالوج_الموحد_2026-10-02.xlsx", read_only=True)
    sheets2 = wb2.sheetnames
    check("جدول الكتالوج الموحد = 7 أوراق", 7, len(sheets2) == 7, {"sheets": sheets2})
    wb2.close()
except Exception as e:
    check("جدول الكتالوج الموحد = 7 أوراق", 7, False, f"EXC {e}")

# ---------- 6. Claim: negotiation guide docx = 13 sections ----------
try:
    import docx
    d = docx.Document(ROOT / "download/دليل_حملة_التفاوض_الافتتاحية_2026-10-02.docx")
    h1 = [p.text for p in d.paragraphs if p.style.name in ("Heading 1", "Title") and p.text.strip()]
    check("دليل حملة التفاوض = 13 قسمًا (H1)", 13, len(h1) == 13, {"h1_count": len(h1), "titles": h1[:15]})
except Exception as e:
    check("دليل حملة التفاوض = 13 قسمًا", 13, False, f"EXC {e}")

# ---------- 7. Claim: gateway pricing = 15 models ----------
try:
    gp = json.loads((ROOT / "research/sv_master_gateway_pricing_20261002.json").read_text())
    n_models = len(gp.get("data", []))
    check("بوابة teamsoclo = 15 موديلًا (وثيقة 02/10)", 15, n_models == 15, {"models": n_models})
except Exception as e:
    check("بوابة teamsoclo = 15 موديلًا", 15, False, f"EXC {e}")

# ---------- 8. Claim: worklog documents SV-NEG-CAMPAIGN (v1.5 basis) ----------
wl = (ROOT / "worklog.md").read_text(encoding="utf-8")
check("worklog يوثق SV-NEG-CAMPAIGN (أساس v1.5)", True, "SV-NEG-CAMPAIGN" in wl, "grep")
check("worklog يوثق G-1a-R1 (أساس v1.4)", True, "Task ID: G-1a-R1" in wl, "grep")
check("worklog يوثق G-1-R3-NEWKEYS (أساس v1.3)", True, "G-1-R3-NEWKEYS" in wl, "grep")
check("worklog يوثق SV-MASTER-CATALOG", True, "SV-MASTER-CATALOG" in wl, "grep")

# ---------- 9. PMRF v1.5 internal consistency ----------
pmrf = (ROOT / "download/PROJECT MASTER REFERENCE FILE.md").read_text(encoding="utf-8")
check("PMRF الحالي يحمل ترويسة v1.5", True, "PMRF v1.5" in pmrf and "v1.5 (2026-10-02T11:05+03:00)" in pmrf, "header")
# NOTE: §1 Identity table says "الإصدار v1.4" while header says v1.5 — known v1.5 defect?
check("جدول الهوية §1 متسق مع الترويسة (v1.5)", True,
      bool(re.search(r"الإصدار\s*\|\s*\*\*v1\.5\*\*|الإصدار.*v1\.5", pmrf)) or "v1.5" in pmrf.split("## 2.")[0].split("الإصدار")[-1][:80] if "الإصدار" in pmrf.split("## 2.")[0] else False,
      "§1 version field vs v1.5 header")

OUT.write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps({
    "timestamp": TS,
    "claim_results": [{ "claim": c["claim"], "result": c["result"], "detail": c["detail"] if c["result"] == "FAIL" else None} for c in report["claim_checks"]],
    "runtime_counts": report["runtime_counts"],
}, ensure_ascii=False, indent=1))
