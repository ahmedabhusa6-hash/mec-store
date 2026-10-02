#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SV-MASTER-CATALOG — Build the unified organized Excel workbook.
Input:  research/sv_master_catalog_consolidated_20261002.json
Output: download/جدول_الكتالوج_الموحد_2026-10-02.xlsx
Sheets: الملخص · StackVault_الحي · StackVault_الأرشيف · الكيانات_الخارجية ·
        مقارنة_الأسعار · بوابة_teamsoclo · المراجعة
"""
import sys, os, json

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import *  # design tokens + style factories
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

SRC = "/home/z/my-project/research/sv_master_catalog_consolidated_20261002.json"
OUT = "/home/z/my-project/download/جدول_الكتالوج_الموحد_2026-10-02.xlsx"

D = json.load(open(SRC, encoding="utf-8"))
E = D["entities"]
AUD = D["audit"]

wb = Workbook()
wb.properties.creator = "Z.ai"

FMT_USD = FORMATS["currency_usd"]
FMT_INT = FORMATS["integer"]
FMT_PCT = FORMATS["percentage"]
FMT_R3 = "0.000"

def wcell(ws, r, c, v, fmt=None, numeric=False):
    cell = ws.cell(row=r, column=c, value=v)
    if fmt: cell.number_format = fmt
    if numeric: cell.alignment = align_number()
    return cell

def build_table(ws, title, headers, rows, numeric_cols=None, fmt_map=None, widths=None, freeze=True):
    """Standard sheet: title B2, headers row4, data rows5+, auto-fit."""
    numeric_cols = numeric_cols or {}
    fmt_map = fmt_map or {}
    last_col = len(headers) + 1
    setup_sheet(ws, title=title, last_col=last_col)
    for ci, h in enumerate(headers, start=2):
        ws.cell(row=4, column=ci, value=h)
    style_header_row(ws, 4, 2, last_col)
    for i, row in enumerate(rows):
        rn = 5 + i
        for ci, v in enumerate(row, start=2):
            fmt = fmt_map.get(ci)
            wcell(ws, rn, ci, v, fmt=fmt, numeric=(ci in numeric_cols))
        style_data_row(ws, rn, 2, last_col, i)
        # re-apply numeric alignment after style pass
        for ci in numeric_cols:
            ws.cell(row=rn, column=ci).alignment = align_number()
            if ci in fmt_map:
                ws.cell(row=rn, column=ci).number_format = fmt_map[ci]
    auto_fit_columns(ws, min_width=8, max_width=30, header_row=4, data_start_row=5)
    if widths:
        for col_letter, w in widths.items():
            ws.column_dimensions[col_letter].width = w
    if freeze:
        ws.freeze_panes = "C5"
    return 5 + len(rows) - 1  # last data row

# ═══ Sheet 1: الملخص ═════════════════════════════════════════════════════
ws = wb.active
ws.title = "الملخص"
rows = []
order = ["stackvault_live", "stackvault_archive", "prodseller", "canboso_premikey",
         "lahastore", "aiversex", "geminishop", "acczone", "digitalcore", "aixpress", "richai"]
for k in order:
    e, a = E[k], AUD["per_entity"][k]
    rows.append([e["display"], e["platform"], a["n"], a["price_min"], a["price_max"],
                 a["in_stock_n"], e["pulled_at"], e["verdict"]])
gw = D["gateway"]
rows.append(["بوابة teamsoclo (نماذج API)", gw["platform"], len(gw["models"]), "", "", "",
             "2026-10-02 02:49 +03", "بوابة الجملة العليا المترابطة مع cb_ (Team Sóc Lọ)"])
last = build_table(
    ws, "الجرد الشامل — كل عمليات سحب المنتجات (حتى 2026-10-02)",
    ["الكيان", "المنصة / النطاق", "عدد المنتجات", "أدنى سعر $", "أعلى سعر $",
     "المتوفر", "تاريخ السحب", "الحكم"],
    rows,
    numeric_cols={4, 5, 6, 7}, fmt_map={4: FMT_USD, 5: FMT_USD},
    widths={"C": 34, "J": 46}, freeze=False)
# totals row
tr = last + 1
wcell(ws, tr, 2, "الإجمالي — سجلات المنتجات (11 كتالوجًا)")
wcell(ws, tr, 4, AUD["total_product_records"], fmt=FMT_INT, numeric=True)
style_total_row(ws, tr, 2, 9)
ws.cell(row=tr, column=4).alignment = align_number()
ws.cell(row=tr, column=4).number_format = FMT_INT
# caption notes
note_r = tr + 2
notes = [
    "المصدر الموحد: research/sv_master_catalog_consolidated_20261002.json · الإصدار المرجعي: PMRF v1.4",
    D["currency_note"],
    "أرشيف StackVault يتضمن costPrice المسرب (أُغلق لاحقًا) — الهوامش: وسيط +20.1% · نطاق +15% إلى +154.3%",
    "المستبعد كهوية cb_ (صفر تطابق): laha · aivx · gshop · acczone · dcore · aixpress · richai · canboso (قاعدة منفصلة)",
    "المتبقي المفتوح: G-7 شراء اختباري (الهوية القانونية لمشغّل cb_) — كل مفاتيح Ahmed تعمل لحساب TG 7334478984",
]
for i, t in enumerate(notes):
    c = ws.cell(row=note_r + i, column=2, value=t)
    c.font = font_caption()

# ═══ Sheet 2: StackVault_الحي (282) ══════════════════════════════════════
ws2 = wb.create_sheet("StackVault_الحي")
live = E["stackvault_live"]["products"]
rows = [[p["id"], p["name"], p["category"], p["prefix"], p["price_usd"],
         p["stock"], "نعم" if p["in_stock"] else "لا", p["description_head"]] for p in live]
build_table(
    ws2, "StackVault — الكتالوج الحي (282 منتجًا · سحب 2026-10-02)",
    ["المعرف", "المنتج", "الفئة", "البادئة", "السعر $", "المخزون", "متوفر", "مقتطف الوصف"],
    rows, numeric_cols={6, 7}, fmt_map={6: FMT_USD, 7: FMT_INT},
    widths={"C": 44, "I": 48})

# ═══ Sheet 3: StackVault_الأرشيف (357 + margin formula) ══════════════════
ws3 = wb.create_sheet("StackVault_الأرشيف")
arch = E["stackvault_archive"]["products"]
rows = [[p["id"], p["name"], p["category"], p["prefix"],
         p["cost_usd"] if p["cost_usd"] is not None else "", p["price_usd"], None, p["stock"]]
        for p in arch]
last = build_table(
    ws3, "StackVault — أرشيف ما قبل الهجرة (357 منتجًا مع costPrice · 29-30/09)",
    ["المعرف", "المنتج", "الفئة", "البادئة", "التكلفة $ (مسربة)", "السعر $", "هامش الربح %", "المخزون"],
    rows, numeric_cols={6, 7, 8, 9}, fmt_map={6: FMT_USD, 7: FMT_USD, 8: FMT_PCT, 9: FMT_INT},
    widths={"C": 44})
# live margin formulas (col H = 8)
for r in range(5, last + 1):
    ws3.cell(row=r, column=8, value=f'=IF(OR(F{r}="",F{r}=0),"",(G{r}-F{r})/F{r})')
    ws3.cell(row=r, column=8).number_format = FMT_PCT
    ws3.cell(row=r, column=8).alignment = align_number()
# conditional color scale on margin column H
ws3.conditional_formatting.add(
    f"H5:H{last}",
    ColorScaleRule(start_type="min", start_color="F8696B",
                   mid_type="percentile", mid_value=50, mid_color="FFEB84",
                   end_type="max", end_color="63BE7B"))

# ═══ Sheet 4: الكيانات_الخارجية (473) ═══════════════════════════════════
ws4 = wb.create_sheet("الكيانات_الخارجية")
ext_order = ["prodseller", "lahastore", "aiversex", "geminishop", "acczone",
             "canboso_premikey", "digitalcore", "aixpress", "richai"]
rows = []
for k in ext_order:
    e = E[k]
    short = e["display"].split(" — ")[0]
    for p in e["products"]:
        note = ""
        if k == "lahastore":
            note = f"الأصلي: {p.get('price_vnd'):,.0f} VND" if p.get("price_vnd") else ""
        elif k == "digitalcore":
            note = f"من ${p['price_from']} بكميات" if p.get("price_from") else ""
        elif k == "richai":
            note = f"سعر وحدتك: ${p['unit_price']}" if p.get("unit_price") else ""
        elif k == "canboso_premikey":
            note = f"بيعت: {p['sold']}" if p.get("sold") is not None else ""
        elif k == "prodseller":
            note = f"publicPrice: ${p['public_price']} · بيع: {p.get('sold')}" if p.get("public_price") else ""
        stock = p.get("stock")
        rows.append([short, p["id"], p["name"], p.get("category") or "", p.get("price_usd"),
                     stock, "نعم" if p.get("in_stock") else "لا", note])
build_table(
    ws4, "الكيانات الخارجية — الكتالوجات المسحوبة بالمفاتيح (473 سجلًا · 9 كيانات)",
    ["الكيان", "المعرف", "المنتج", "النوع/الفئة", "السعر $", "المخزون", "متوفر", "ملاحظة"],
    rows, numeric_cols={6, 7}, fmt_map={6: FMT_USD, 7: FMT_INT},
    widths={"C": 30, "D": 42, "I": 34})

# ═══ Sheet 5: مقارنة_الأسعار (111 pairs) ═════════════════════════════════
ws5 = wb.create_sheet("مقارنة_الأسعار")
pairs = D["cross_matches"]["cb_vs_canboso"]["pairs"]
rows = [[p["name"], p["cb_price"], p["canboso_price"], None, ""] for p in pairs]
last = build_table(
    ws5, "مقارنة الأسعار — cb_ مقابل canboso/PremiKey (111 زوجًا متطابقًا بالاسم)",
    ["المنتج", "سعر cb_ $", "سعر canboso $", "النسبة cb_/canboso", "قراءة"],
    rows, numeric_cols={3, 4, 5}, fmt_map={3: FMT_USD, 4: FMT_USD, 5: FMT_R3},
    widths={"C": 46})
for i, r in enumerate(range(5, last + 1)):
    ws5.cell(row=r, column=5, value=f'=IFERROR(C{r}/D{r},"")')
    ws5.cell(row=r, column=5).number_format = FMT_R3
    ws5.cell(row=r, column=5).alignment = align_number()
    ws5.cell(row=r, column=6, value=f'=IF(E{r}="","",IF(E{r}<0.8,"cb_ أرخص بوضوح",IF(E{r}<=1.05,"تقارب","canboso أرخص")))')
# reversed color scale: low ratio (cb_ cheaper) = green
ws5.conditional_formatting.add(
    f"E5:E{last}",
    ColorScaleRule(start_type="min", start_color="63BE7B",
                   mid_type="percentile", mid_value=50, mid_color="FFEB84",
                   end_type="max", end_color="F8696B"))
tr = last + 1
wcell(ws5, tr, 2, "الوسيط")
wcell(ws5, tr, 5, AUD["cb_canboso_ratio"]["median"], fmt=FMT_R3, numeric=True)
style_total_row(ws5, tr, 2, 6)
ws5.cell(row=tr, column=5).alignment = align_number()
ws5.cell(row=tr, column=5).number_format = FMT_R3
c = ws5.cell(row=tr + 2, column=2,
             value=f"وسيط النسبة {AUD['cb_canboso_ratio']['median']} (p25 {AUD['cb_canboso_ratio']['p25']} · p75 {AUD['cb_canboso_ratio']['p75']}) — cb_ أرخص إجمالًا → canboso ليس المنبع الظاهر · الاستثناء: خط Deepseek API (أرخص على canboso 1.6-1.7×)")
c.font = font_caption()

# ═══ Sheet 6: بوابة_teamsoclo (15) ═══════════════════════════════════════
ws6 = wb.create_sheet("بوابة_teamsoclo")
models = D["gateway"]["models"]
rows = [[m["model"], m["quota_type"], m["model_ratio"], m["model_price"], m["completion_ratio"]]
        for m in models]
build_table(
    ws6, f"بوابة teamsoclo — تسعيرة النماذج ({len(models)} موديلًا · الإصدار {D['gateway']['pricing_version'][:8]})",
    ["الموديل", "نوع الحصة", "model_ratio", "model_price", "completion_ratio"],
    rows, numeric_cols={4, 5, 6}, fmt_map={4: FMT_R3, 5: FMT_R3, 6: FMT_R3},
    widths={"C": 26})

# ═══ Sheet 7: المراجعة (Review) ══════════════════════════════════════════
ws7 = wb.create_sheet("المراجعة")
ws7.sheet_properties.tabColor = "FFC000"
checks = [
    ["فحص", "المتوقع", "الفعلي", "الحالة"],
    ["منتجات StackVault الحي", 282, "=COUNTA('StackVault_الحي'!B5:B286)", None],
    ["منتجات StackVault الأرشيف", 357, "=COUNTA('StackVault_الأرشيف'!B5:B361)", None],
    ["سجلات الكيانات الخارجية", 473, "=COUNTA('الكيانات_الخارجية'!B5:B477)", None],
    ["أزواج مقارنة الأسعار", 111, "=COUNTA('مقارنة_الأسعار'!B5:B115)", None],
    ["موديلات البوابة", 15, "=COUNTA('بوابة_teamsoclo'!B5:B19)", None],
    ["إجمالي سجلات المنتجات", 1112, "=D5+D6+D7", None],
]
setup_sheet(ws7, title="المراجعة — فحوص التطابق", last_col=5)
for ci, h in enumerate(checks[0], start=2):
    ws7.cell(row=4, column=ci, value=h)
style_header_row(ws7, 4, 2, 5)
for i, row in enumerate(checks[1:]):
    rn = 5 + i
    ws7.cell(row=rn, column=2, value=row[0])
    wcell(ws7, rn, 3, row[1], fmt=FMT_INT, numeric=True)
    ws7.cell(row=rn, column=4, value=row[2])
    ws7.cell(row=rn, column=5, value=f'=IF(C{rn}=D{rn},"✓ مطابق","✗ راجع")')
    style_data_row(ws7, rn, 2, 5, i)
    ws7.cell(row=rn, column=3).alignment = align_number()
    ws7.cell(row=rn, column=4).alignment = align_number()
auto_fit_columns(ws7, min_width=10, max_width=34, header_row=4, data_start_row=5)

wb.save(OUT)
print(f"OK → {OUT} ({os.path.getsize(OUT)//1024}KB)")
print("Sheets:", wb.sheetnames)
